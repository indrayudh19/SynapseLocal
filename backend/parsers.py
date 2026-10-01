"""
backend/parsers.py
Module 1.1: Multi-Modal Document Parsing

Extracts structured text and metadata from PDF, TXT, MD, and PPTX files.
Preserves heading hierarchy for contextual chunking downstream.
"""
import io
import re
from dataclasses import dataclass, field
from typing import List


@dataclass
class ParsedSection:
    """A single section of a parsed document with heading context."""
    heading: str
    content: str
    source_file: str
    page_or_slide: int = 0


def parse_pdf(file_bytes: bytes, filename: str) -> List[ParsedSection]:
    """Extract text from PDF, grouping by pages."""
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise ImportError("PyPDF2 is required for PDF parsing. Install: pip install PyPDF2")

    reader = PdfReader(io.BytesIO(file_bytes))
    sections = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            sections.append(ParsedSection(
                heading=f"Page {i + 1}",
                content=text,
                source_file=filename,
                page_or_slide=i + 1
            ))
    return sections


def parse_txt(file_bytes: bytes, filename: str) -> List[ParsedSection]:
    """Extract plain text, splitting by double-newline paragraphs."""
    text = file_bytes.decode("utf-8", errors="replace")
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    sections = []
    for i, para in enumerate(paragraphs):
        sections.append(ParsedSection(
            heading=f"Section {i + 1}",
            content=para,
            source_file=filename,
            page_or_slide=i + 1
        ))
    return sections


def parse_markdown(file_bytes: bytes, filename: str) -> List[ParsedSection]:
    """Extract markdown text, splitting by heading hierarchy (# / ## / ###)."""
    text = file_bytes.decode("utf-8", errors="replace")
    # Split on heading markers while keeping the heading line
    parts = re.split(r"(?=^#{1,6}\s)", text, flags=re.MULTILINE)
    sections = []
    for i, part in enumerate(parts):
        part = part.strip()
        if not part:
            continue
        # Extract heading from first line if it starts with #
        lines = part.split("\n", 1)
        heading_match = re.match(r"^(#{1,6})\s+(.*)", lines[0])
        if heading_match:
            heading = heading_match.group(2).strip()
            content = lines[1].strip() if len(lines) > 1 else ""
        else:
            heading = f"Section {i + 1}"
            content = part
        if content:
            sections.append(ParsedSection(
                heading=heading,
                content=content,
                source_file=filename,
                page_or_slide=i + 1
            ))
    return sections


def parse_pptx(file_bytes: bytes, filename: str) -> List[ParsedSection]:
    """Extract slide titles, body text, and speaker notes from PPTX."""
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx is required for PPTX parsing. Install: pip install python-pptx")

    prs = Presentation(io.BytesIO(file_bytes))
    sections = []
    for i, slide in enumerate(prs.slides):
        title = ""
        body_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if shape == slide.shapes.title:
                    title = text
                elif text:
                    body_parts.append(text)
        # Speaker notes
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        
        combined = "\n".join(body_parts)
        if notes:
            combined += f"\n[Speaker Notes]: {notes}"
        
        if combined.strip():
            sections.append(ParsedSection(
                heading=title or f"Slide {i + 1}",
                content=combined,
                source_file=filename,
                page_or_slide=i + 1
            ))
    return sections


def parse_document(file_bytes: bytes, filename: str) -> List[ParsedSection]:
    """Route to the correct parser based on file extension."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    
    parsers = {
        "pdf": parse_pdf,
        "txt": parse_txt,
        "md": parse_markdown,
        "pptx": parse_pptx,
    }
    
    parser_fn = parsers.get(ext)
    if parser_fn is None:
        raise ValueError(f"Unsupported file type: .{ext}. Supported: {list(parsers.keys())}")
    
    return parser_fn(file_bytes, filename)
